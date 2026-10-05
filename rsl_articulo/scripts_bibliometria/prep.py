# -*- coding: utf-8 -*-
import re, json, collections
from openpyxl import load_workbook
BASE='../tex/Rslunlearnable-main/Article_Title/articulos-compilacion/bibliometria/bibtex_maestro (2).bib'
m=open(BASE,encoding='utf8',errors='ignore').read()
ents=re.split(r'\n(?=@\w+\{)',m)[1:]
def fld(e,f):
    r=re.search(r'^\s*'+f+r'\s*=\s*\{(.*?)\},?\s*$',e,flags=re.M|re.S|re.I)
    return re.sub(r'\s+',' ',r.group(1)).strip() if r else None
B={}
for e in ents:
    d=(fld(e,'doi') or '').lower().strip(); t=re.sub(r'[^a-z0-9]','',(fld(e,'title') or '').lower())
    B[d or t]=dict(title=fld(e,'title'),year=fld(e,'year'),akw=fld(e,'author_keywords'),ikw=fld(e,'keywords'),ab=fld(e,'abstract'),db=fld(e,'db_source'))
wb=load_workbook('../x/sabana_RSL_v4.7_2026-10-05.xlsx'); x9=wb['X9 Cribado']; art=wb['ARTICULOS SELECCIONADOS']
INC={art.cell(r,14).value:art.cell(r,1).value for r in range(6,41)}
dom={art.cell(r,14).value:art.cell(r,12).value for r in range(6,41)}
recs=[]
for r in range(5,348):
    if not x9.cell(r,1).value: continue
    doi=(x9.cell(r,2).value or '').lower().strip(); t=re.sub(r'[^a-z0-9]','',(x9.cell(r,3).value or '').lower())
    k=doi if doi in B else t
    b=B[k]; oid=x9.cell(r,9).value
    cons=x9.cell(r,15).value; code=x9.cell(r,16).value
    recs.append(dict(id=x9.cell(r,1).value,oid=oid,title=x9.cell(r,3).value,year=int(x9.cell(r,4).value),cons=cons,code=code,akw=b['akw'],ikw=b['ikw'],ab=b['ab'],inc=oid in INC,dom=dom.get(oid)))
TEM=[x for x in recs if x['cons']=='Incluir' or x['code'] in ('EX1','EX2','EX3','EX5b','EX8','EX9')]
print('tematico',len(TEM),'incluidos en bases',sum(1 for x in recs if x['inc']))
json.dump(dict(recs=recs,tem=[x['id'] for x in TEM]),open('recs.json','w'),ensure_ascii=False)
yrs=collections.Counter(x['year'] for x in TEM); print(sorted(yrs.items()))
print(collections.Counter(x['code'] or 'Incluir' for x in TEM))
